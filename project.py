import pytesseract
from PIL import Image
from docx import Document
import re


def main():
    image_name = input("Enter identity card image filename: ").strip()
    template_name = input("Enter contract template filename: ").strip()

    text = extract_text_from_image(image_name)
    info = parse_identity_card(text)
    validate_info(info)
    fill_contract(template_name, info, "output_contract.docx")

    print("Extracted data:", info)
    print("Contract generated: output_contract.docx")


def extract_text_from_image(image_name):
    """
    Extract text from image using Tesseract (English mode).
    """

    try:
        img = Image.open(image_name).convert("L")
        text = pytesseract.image_to_string(img, lang="eng")
        return text
    except FileNotFoundError:
        raise FileNotFoundError("Image file not found.")


def parse_identity_card(text):
    """
    Parse structured English mock identity card.
    Expected format:
    Key: Value
    """

    info = {}
    lines = text.split("\n")

    for line in lines:
        if ":" in line:
            key, value = line.split(":", 1)
            info[key.strip()] = value.strip()

    return {
        "id": info.get("Personal ID", ""),
        "name": info.get("Full name", ""),
        "dob": info.get("Date of birth", ""),
        "gender": info.get("Sex", ""),
        "nationality": info.get("Nationality", ""),
        "address": info.get("Place of residence", ""),
        "issue_date": info.get("Date of issue", ""),
        "issue_place": info.get("Place of issue", "")
    }


def validate_info(info):
    """
    Validate extracted fields.
    """

    if not re.fullmatch(r"\d{12}", info["id"]):
        raise ValueError("Invalid Personal ID (must be 12 digits).")

    if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", info["dob"]):
        raise ValueError("Invalid Date of Birth format (dd/mm/yyyy).")

    if not info["name"].isupper():
        raise ValueError("Name must be uppercase.")


def fill_contract(template_name, info, output_name):
    """
    Replace placeholders in contract template.
    """

    try:
        doc = Document(template_name)

        for p in doc.paragraphs:
            p.text = p.text.replace("{{NAME}}", info["name"])
            p.text = p.text.replace("{{DOB}}", info["dob"])
            p.text = p.text.replace("{{ID}}", info["id"])
            p.text = p.text.replace("{{GENDER}}", info["gender"])
            p.text = p.text.replace("{{ADDRESS}}", info["address"])
            p.text = p.text.replace("{{ISSUE_DATE}}", info["issue_date"])
            p.text = p.text.replace("{{ISSUE_PLACE}}", info["issue_place"])

        doc.save(output_name)

    except FileNotFoundError:
        raise FileNotFoundError("Template file not found.")


if __name__ == "__main__":
    main()
